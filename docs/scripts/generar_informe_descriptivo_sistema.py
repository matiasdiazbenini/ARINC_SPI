from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUT_DIR = DOCS / "entrega_tutor"
ASSET_DIR = OUT_DIR / "recursos_informe"
DOCX_PATH = OUT_DIR / "informe_descriptivo_sistema_arinc429.docx"

FWD_SINGLE = ROOT / "imagenes" / "FWD" / "GP2 (CH1)" / "TEK0000.JPG"
FWD_DIFF = ROOT / "imagenes" / "FWD" / "CH1 - CH2" / "TEK0011.JPG"
REV_SINGLE = ROOT / "imagenes" / "REV" / "GP4 (CH1)" / "TEK0012.JPG"
REV_DIFF = ROOT / "imagenes" / "REV" / "CH1 - CH2" / "TEK0019.JPG"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

NAVY = "17324D"
BLUE = "2E74B5"
CYAN = "2D8C9F"
GREEN = "2F7D5A"
GOLD = "B7791F"
RED = "A33A3A"
INK = "1E2933"
MUTED = "5D6975"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E8F3ED"
PALE_GOLD = "FFF4DC"
WHITE = "FFFFFF"
BLACK = "000000"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def font_path(bold: bool = False) -> str:
    candidate = Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf")
    if candidate.exists():
        return str(candidate)
    fallback = Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf")
    return str(fallback)


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path(bold), size=size)


def wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str = INK,
    line_gap: int = 8,
) -> None:
    x1, y1, x2, y2 = box
    lines: list[str] = []
    for paragraph in text.split("\n"):
        if paragraph.strip():
            lines.extend(wrap(draw, paragraph.strip(), font, max(20, x2 - x1 - 28)))
        else:
            lines.append("")
    sample_height = draw.textbbox((0, 0), "Ag", font=font)[3]
    heights = [draw.textbbox((0, 0), line, font=font)[3] if line else sample_height for line in lines]
    total_h = sum(heights) + line_gap * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, h in zip(lines, heights):
        if line:
            width = draw.textbbox((0, 0), line, font=font)[2]
            draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=font, fill=f"#{fill}")
        y += h + line_gap


def draw_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    detail: str = "",
    fill: str = WHITE,
    outline: str = BLUE,
    title_color: str = NAVY,
) -> None:
    draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{outline}", width=4)
    x1, y1, x2, y2 = box
    title_box = (x1 + 18, y1 + 12, x2 - 18, y1 + 64)
    draw_centered_text(draw, title_box, title, pil_font(28, True), title_color)
    if detail:
        detail_box = (x1 + 22, y1 + 64, x2 - 22, y2 - 16)
        draw_centered_text(draw, detail_box, detail, pil_font(22), INK, 5)


def draw_arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    color: str = BLUE,
    width: int = 6,
    label: str | None = None,
    label_offset: tuple[int, int] = (0, -32),
    dashed: bool = False,
) -> None:
    x1, y1 = start
    x2, y2 = end
    if dashed:
        segments = 14
        for i in range(segments):
            if i % 2 == 0:
                a = i / segments
                b = (i + 1) / segments
                draw.line(
                    (x1 + (x2 - x1) * a, y1 + (y2 - y1) * a,
                     x1 + (x2 - x1) * b, y1 + (y2 - y1) * b),
                    fill=f"#{color}",
                    width=width,
                )
    else:
        draw.line((x1, y1, x2, y2), fill=f"#{color}", width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    head = 22
    points = [
        (x2, y2),
        (x2 - head * math.cos(angle - math.pi / 6), y2 - head * math.sin(angle - math.pi / 6)),
        (x2 - head * math.cos(angle + math.pi / 6), y2 - head * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(points, fill=f"#{color}")
    if label:
        font = pil_font(21, True)
        tx = (x1 + x2) / 2 + label_offset[0]
        ty = (y1 + y2) / 2 + label_offset[1]
        bbox = draw.textbbox((0, 0), label, font=font)
        pad = 8
        draw.rounded_rectangle(
            (tx - pad, ty - pad, tx + bbox[2] + pad, ty + bbox[3] + pad),
            radius=7,
            fill="#FFFFFF",
        )
        draw.text((tx, ty), label, font=font, fill=f"#{color}")


def draw_label(
    draw: ImageDraw.ImageDraw,
    center: tuple[int, int],
    text: str,
    *,
    color: str = INK,
    fill: str = WHITE,
    outline: str | None = None,
    font_size: int = 21,
    bold: bool = True,
    padding: int = 9,
) -> None:
    font = pil_font(font_size, bold)
    bbox = draw.textbbox((0, 0), text, font=font)
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    x = center[0] - width / 2
    y = center[1] - height / 2
    draw.rounded_rectangle(
        (x - padding, y - padding, x + width + padding, y + height + padding),
        radius=8,
        fill=f"#{fill}",
        outline=f"#{outline}" if outline else None,
        width=2 if outline else 1,
    )
    draw.text((x, y), text, font=font, fill=f"#{color}")


def draw_arrowhead(
    draw: ImageDraw.ImageDraw,
    previous: tuple[int, int],
    endpoint: tuple[int, int],
    color: str,
    size: int = 22,
) -> None:
    angle = math.atan2(endpoint[1] - previous[1], endpoint[0] - previous[0])
    points = [
        endpoint,
        (
            endpoint[0] - size * math.cos(angle - math.pi / 6),
            endpoint[1] - size * math.sin(angle - math.pi / 6),
        ),
        (
            endpoint[0] - size * math.cos(angle + math.pi / 6),
            endpoint[1] - size * math.sin(angle + math.pi / 6),
        ),
    ]
    draw.polygon(points, fill=f"#{color}")


def draw_polyline_arrow(
    draw: ImageDraw.ImageDraw,
    points: list[tuple[int, int]],
    *,
    color: str = BLUE,
    width: int = 6,
    arrow_end: bool = True,
    arrow_start: bool = False,
) -> None:
    draw.line(points, fill=f"#{color}", width=width, joint="curve")
    if arrow_end and len(points) >= 2:
        draw_arrowhead(draw, points[-2], points[-1], color)
    if arrow_start and len(points) >= 2:
        draw_arrowhead(draw, points[1], points[0], color)


def draw_junction(draw: ImageDraw.ImageDraw, point: tuple[int, int], color: str, radius: int = 9) -> None:
    x, y = point
    draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=f"#{color}", outline="#FFFFFF", width=2)


def new_canvas(width: int, height: int, title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), title, font=pil_font(38, True), fill=f"#{NAVY}")
    if subtitle:
        draw.text((60, 88), subtitle, font=pil_font(23), fill=f"#{MUTED}")
    draw.line((60, 128, width - 60, 128), fill=f"#{PALE_BLUE}", width=4)
    return image, draw


def save_diagram(image: Image.Image, name: str) -> Path:
    path = ASSET_DIR / name
    image.save(path, quality=95)
    return path


def create_architecture_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        1040,
        "Arquitectura completa",
        "Flujo de datos desde la generación lógica hasta la observabilidad",
    )
    tx = (80, 200, 390, 410)
    rx = (910, 200, 1220, 410)
    sniffer = (500, 620, 830, 880)
    host = (990, 620, 1320, 880)
    notebook = (1450, 600, 1740, 900)

    draw_box(draw, tx, "ARINC-TX", "Pico\nGenera palabras FWD\n100 kbps", PALE_BLUE, BLUE)
    draw_box(draw, rx, "ARINC-RX", "Pico\nRecibe FWD\nModo stream sin ACK", PALE_GREEN, GREEN)
    draw_box(draw, sniffer, "ARINC-SNIFFER", "Pico\nObservación pasiva\nParidad + filtro\nSnapshot SPI", PALE_GOLD, GOLD)
    draw_box(draw, host, "Raspberry Pi 3B+", "SPI master\nBridge HTTP local\nDashboard Flask", LIGHT, NAVY)
    draw_box(draw, notebook, "Notebook", "Ethernet directo\nDashboard\nPrometheus\nGrafana", PALE_BLUE, CYAN)

    # FWD is the cable between TX and RX. The SNIFFER taps the same conductors in parallel.
    fwd_a_y, fwd_b_y = 275, 345
    draw.line((tx[2], fwd_a_y, rx[0], fwd_a_y), fill=f"#{BLUE}", width=8)
    draw.line((tx[2], fwd_b_y, rx[0], fwd_b_y), fill=f"#{CYAN}", width=8)
    draw_arrow(draw, (500, 235), (800, 235), NAVY, 5)
    draw_label(draw, (650, 190), "FWD: par de conductores GP2 / GP3", color=NAVY, font_size=21)

    tap_a = (610, fwd_a_y)
    tap_b = (700, fwd_b_y)
    draw_junction(draw, tap_a, BLUE)
    draw_junction(draw, tap_b, CYAN)
    draw_polyline_arrow(draw, [tap_a, (610, 550), (600, 550), (600, sniffer[1])], color=BLUE, width=6)
    draw_polyline_arrow(draw, [tap_b, (700, 520), (730, 520), (730, sniffer[1])], color=CYAN, width=6)
    draw_label(draw, (1080, 500), "Derivación pasiva en paralelo sobre el cable", color=GOLD, outline=GOLD, font_size=20)

    # REV is a separate optional simplex cable, only observed in the final stream architecture.
    draw.rounded_rectangle((80, 500, 350, 590), radius=14, fill="#FFFFFF", outline=f"#{GOLD}", width=3)
    draw_centered_text(draw, (95, 510, 335, 580), "Segundo canal simplex\nREV opcional", pil_font(20, True), GOLD, 4)
    draw.line((350, 530, 460, 530, 460, 735, sniffer[0], 735), fill=f"#{GOLD}", width=5)
    draw.line((350, 570, 430, 570, 430, 785, sniffer[0], 785), fill=f"#{GOLD}", width=5)
    draw_label(draw, (255, 650), "solo observado", color=GOLD, font_size=18)

    draw_polyline_arrow(draw, [(sniffer[2], 750), (host[0], 750)], color=NAVY, width=8, arrow_start=True)
    draw_label(draw, ((sniffer[2] + host[0]) // 2, 705), "SPI PIO-frame · 8 MHz", color=NAVY, font_size=20)
    draw_polyline_arrow(draw, [(host[2], 750), (notebook[0], 750)], color=CYAN, width=8, arrow_start=True)
    draw_label(draw, ((host[2] + notebook[0]) // 2, 705), "Ethernet · HTTP / métricas", color=CYAN, font_size=19)
    return save_diagram(image, "01_arquitectura_completa.png")


def create_wiring_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        1160,
        "Conexionado principal",
        "Señales lógicas ARINC, enlace SPI y red Ethernet directa",
    )
    tx = (70, 180, 360, 400)
    rx = (950, 180, 1240, 400)
    sniffer = (500, 590, 830, 900)
    host = (1030, 590, 1360, 900)
    notebook = (1510, 630, 1740, 860)

    draw_box(draw, tx, "ARINC-TX", "GP2 = FWD_A\nGP3 = FWD_B\nGND común", PALE_BLUE, BLUE)
    draw_box(draw, rx, "ARINC-RX", "GP2 = FWD_A\nGP3 = FWD_B\nREV inactivo\nGND común", PALE_GREEN, GREEN)
    draw_box(draw, sniffer, "ARINC-SNIFFER", "GP2/GP3 = FWD\nGP4/GP5 = REV\nGP16 = MOSI\nGP17 = CSn\nGP18 = SCK\nGP19 = MISO", PALE_GOLD, GOLD)
    draw_box(draw, host, "Raspberry Pi 3B+", "GPIO10 = MOSI\nGPIO8 = CS manual\nGPIO11 = SCLK\nGPIO9 = MISO\neth0 = 192.168.50.2", LIGHT, NAVY)
    draw_box(draw, notebook, "Notebook", "Ethernet\n192.168.50.1/24", PALE_BLUE, CYAN)

    # Point-to-point FWD cable. The tap is on the cable, not on either board.
    fwd_a_y, fwd_b_y = 260, 335
    draw.line((tx[2], fwd_a_y, rx[0], fwd_a_y), fill=f"#{BLUE}", width=8)
    draw.line((tx[2], fwd_b_y, rx[0], fwd_b_y), fill=f"#{CYAN}", width=8)
    draw_arrow(draw, (555, 220), (755, 220), NAVY, 5)
    draw_label(draw, (655, 180), "Cable FWD: dos conductores", color=NAVY, font_size=22)
    draw_label(draw, (430, 235), "FWD_A", color=BLUE, font_size=18)
    draw_label(draw, (430, 360), "FWD_B", color=CYAN, font_size=18)

    tap_a = (620, fwd_a_y)
    tap_b = (710, fwd_b_y)
    draw_junction(draw, tap_a, BLUE)
    draw_junction(draw, tap_b, CYAN)
    draw_polyline_arrow(draw, [tap_a, (620, 505), (590, 505), (590, sniffer[1])], color=BLUE, width=6)
    draw_polyline_arrow(draw, [tap_b, (710, 535), (740, 535), (740, sniffer[1])], color=CYAN, width=6)
    draw_label(draw, (1050, 480), "Toma pasiva en paralelo", color=GOLD, outline=GOLD, font_size=20)

    # Optional REV is another cable pair, also observed in parallel.
    draw_label(draw, (245, 615), "Cable REV opcional", color=GOLD, outline=GOLD, font_size=19)
    draw.line((90, 665, sniffer[0], 665), fill=f"#{GOLD}", width=5)
    draw.line((90, 720, sniffer[0], 720), fill=f"#{GOLD}", width=5)
    draw.text((110, 635), "REV_A", font=pil_font(18, True), fill=f"#{GOLD}")
    draw.text((110, 730), "REV_B", font=pil_font(18, True), fill=f"#{GOLD}")

    # SPI is a separate point-to-point host link.
    spi_y = 745
    draw_polyline_arrow(draw, [(sniffer[2], spi_y), (host[0], spi_y)], color=NAVY, width=9, arrow_start=True)
    draw_label(draw, ((sniffer[2] + host[0]) // 2, spi_y - 48), "SPI: MOSI · MISO · SCK · CSn", color=NAVY, font_size=18)
    draw_polyline_arrow(draw, [(host[2], spi_y), (notebook[0], spi_y)], color=CYAN, width=8, arrow_start=True)
    draw_label(draw, ((host[2] + notebook[0]) // 2, spi_y - 48), "Ethernet", color=CYAN, font_size=19)

    draw.text((85, 940), "Nota eléctrica", font=pil_font(25, True), fill=f"#{NAVY}")
    note = (
        "Las líneas FWD/REV de esta fase son GPIO de 3,3 V con masa común. "
        "No deben conectarse directamente a un bus ARINC 429 de campo."
    )
    draw.rounded_rectangle((80, 985, 1160, 1115), radius=16, fill=f"#{PALE_GOLD}", outline=f"#{GOLD}", width=3)
    draw_centered_text(draw, (105, 995, 1135, 1105), note, pil_font(24), INK, 8)
    return save_diagram(image, "02_conexionado_principal.png")


def create_arinc_model_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        1020,
        "Modelo lógico ARINC 429",
        "Palabra de 32 bits y codificación bipolar con retorno a cero",
    )
    fields = [
        ("Label", 8, BLUE),
        ("SDI", 2, CYAN),
        ("Datos", 19, GREEN),
        ("SSM", 2, GOLD),
        ("P", 1, RED),
    ]
    x = 90
    y1, y2 = 190, 350
    usable = 1620
    for name, bits, color in fields:
        width = usable * bits / 32
        draw.rectangle((x, y1, x + width, y2), fill=f"#{color}", outline="#FFFFFF", width=4)
        draw_centered_text(draw, (int(x), y1, int(x + width), y2), f"{name}\n{bits} bit{'s' if bits > 1 else ''}", pil_font(24, True), WHITE)
        x += width
    draw.text((90, 370), "Bits 1-8", font=pil_font(19), fill=f"#{MUTED}")
    draw.text((1490, 370), "Bit 32", font=pil_font(19), fill=f"#{MUTED}")

    origin_y = 690
    active = 95
    null = 95
    start_x = 180
    bits = [1, 0, 1, 1, 0]
    waveform_end = start_x + len(bits) * (active + null)
    draw.line((120, origin_y, waveform_end + 40, origin_y), fill=f"#{MUTED}", width=4)
    for i, bit in enumerate(bits):
        bx = start_x + i * (active + null)
        level = origin_y - 130 if bit else origin_y + 130
        draw.line((bx, origin_y, bx, level), fill=f"#{BLUE if bit else RED}", width=5)
        draw.line((bx, level, bx + active, level), fill=f"#{BLUE if bit else RED}", width=10)
        draw.line((bx + active, level, bx + active, origin_y), fill=f"#{BLUE if bit else RED}", width=5)
        draw.line((bx + active, origin_y, bx + active + null, origin_y), fill=f"#{MUTED}", width=8)
        draw_centered_text(draw, (bx, 850, bx + active + null, 900), str(bit), pil_font(25, True), INK)

    draw.text((120, 520), "+3,3 V", font=pil_font(22, True), fill=f"#{BLUE}")
    draw.text((120, 670), "0 V", font=pil_font(22, True), fill=f"#{MUTED}")
    draw.text((120, 820), "-3,3 V", font=pil_font(22, True), fill=f"#{RED}")
    draw.text((120, 465), "Nivel diferencial lógico", font=pil_font(24, True), fill=f"#{NAVY}")

    first_bit_end = start_x + active + null
    draw.line((start_x, 935, first_bit_end, 935), fill=f"#{NAVY}", width=4)
    draw_arrowhead(draw, (start_x + 20, 935), (start_x, 935), NAVY, 16)
    draw_arrowhead(draw, (first_bit_end - 20, 935), (first_bit_end, 935), NAVY, 16)
    draw_label(draw, ((start_x + first_bit_end) // 2, 965), "10 µs por bit", color=NAVY, font_size=19)

    info_box = (1260, 500, 1710, 840)
    draw.rounded_rectangle(info_box, radius=18, fill=f"#{LIGHT}", outline=f"#{NAVY}", width=4)
    draw_centered_text(draw, (1290, 520, 1680, 590), "Temporización", pil_font(28, True), NAVY)
    draw_centered_text(
        draw,
        (1300, 600, 1670, 810),
        "Cada bit ocupa 10 µs.\n\n5 µs: nivel activo\n5 µs: nivel NULL\n\nVelocidad: 100 kbps",
        pil_font(24),
        INK,
        8,
    )
    return save_diagram(image, "03_modelo_arinc_logico.png")


def create_internal_tx_rx_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        1050,
        "Flujo interno: ARINC-TX y ARINC-RX",
        "Generación, serialización, FIFO y procesamiento de palabras",
    )
    draw.text((70, 155), "ARINC-TX", font=pil_font(29, True), fill=f"#{BLUE}")
    tx_boxes = [
        ((70, 220, 330, 380), "Generador", "Labels, SDI,\ndatos y SSM"),
        ((400, 220, 660, 380), "Paridad", "Cálculo de\nparidad impar"),
        ((730, 220, 990, 380), "Cola TX", "Palabras de\n32 bits"),
        ((1060, 220, 1320, 380), "PIO TX", "Serializa a\n100 kbps"),
        ((1390, 220, 1720, 380), "GP2 / GP3", "Par lógico FWD\nRZ bipolar"),
    ]
    for box, title, detail in tx_boxes:
        draw_box(draw, box, title, detail, PALE_BLUE, BLUE)
    for left, right in zip(tx_boxes, tx_boxes[1:]):
        draw_arrow(draw, (left[0][2], 300), (right[0][0], 300), BLUE, 6)

    draw.text((70, 500), "ARINC-RX", font=pil_font(29, True), fill=f"#{GREEN}")
    rx_boxes = [
        ((70, 570, 330, 760), "GP2 / GP3", "Muestreo del\npar FWD"),
        ((400, 570, 660, 760), "PIO RX", "Máquina de estado\nreconstruye 32 bits"),
        ((730, 570, 990, 760), "FIFO RX", "Hasta 4 palabras\nsin FIFO_JOIN_RX"),
        ((1060, 570, 1320, 760), "CPU", "Drena FIFO\nValida estructura"),
        ((1390, 570, 1720, 760), "Aplicación", "Contadores,\nvalores y estado"),
    ]
    for box, title, detail in rx_boxes:
        draw_box(draw, box, title, detail, PALE_GREEN, GREEN)
    for left, right in zip(rx_boxes, rx_boxes[1:]):
        draw_arrow(draw, (left[0][2], 665), (right[0][0], 665), GREEN, 6)

    note = (
        "El modo stream evita lotes fijos: TX emite ráfagas de longitud variable y RX drena el FIFO "
        "continuamente. Las operaciones extensas se mantienen fuera del lazo de captura."
    )
    draw.rounded_rectangle((260, 855, 1540, 980), radius=16, fill=f"#{LIGHT}", outline=f"#{NAVY}", width=3)
    draw_centered_text(draw, (285, 870, 1515, 965), note, pil_font(23), INK, 7)
    return save_diagram(image, "04_flujo_interno_tx_rx.png")


def create_internal_sniffer_host_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        1180,
        "Flujo interno: SNIFFER y Raspberry Pi 3B+",
        "Captura pasiva, clasificación, snapshot, SPI y servicios de visualización",
    )
    draw.text((70, 155), "ARINC-SNIFFER", font=pil_font(29, True), fill=f"#{GOLD}")
    sniffer_boxes = [
        ((60, 215, 315, 405), "Entradas", "FWD GP2/GP3\nREV GP4/GP5"),
        ((360, 215, 615, 405), "PIO RX", "Dos máquinas\nde estado"),
        ((660, 215, 915, 405), "FIFO / CPU", "Reconstruye\npalabras"),
        ((960, 215, 1215, 405), "Integridad", "Paridad estricta\nLabel + SDI"),
        ((1260, 215, 1535, 405), "Snapshot", "16 slots\nContadores"),
    ]
    for box, title, detail in sniffer_boxes:
        draw_box(draw, box, title, detail, PALE_GOLD, GOLD)
    for left, right in zip(sniffer_boxes, sniffer_boxes[1:]):
        draw_arrow(draw, (left[0][2], 310), (right[0][0], 310), GOLD, 6)

    # Snapshot data is copied to a dedicated SPI response block through a clear lower lane.
    spi_slave = (60, 500, 335, 665)
    draw_box(draw, spi_slave, "PIO SPI slave", "Solicitud y respuesta\n32 bytes por trama", PALE_BLUE, NAVY)
    draw_polyline_arrow(
        draw,
        [(1397, 405), (1397, 455), (197, 455), (197, spi_slave[1])],
        color=NAVY,
        width=6,
    )
    draw_label(draw, (800, 420), "Copia atómica del snapshot para la respuesta SPI", color=NAVY, font_size=19)

    draw.text((70, 755), "RASPBERRY PI 3B+", font=pil_font(29, True), fill=f"#{NAVY}")
    host_boxes = [
        ((60, 820, 315, 1020), "SPI master", "8 MHz\nCS manual GPIO8"),
        ((360, 820, 615, 1020), "Bridge", "Polling + lock\n127.0.0.1:5100"),
        ((660, 820, 915, 1020), "HTTP/JSON", "stats, control\ny filtros"),
        ((960, 820, 1215, 1020), "Flask", "Dashboard\n0.0.0.0:5000"),
        ((1260, 820, 1535, 1020), "Registro", "CSV + JSON\nPDF estadístico"),
    ]
    for box, title, detail in host_boxes:
        draw_box(draw, box, title, detail, LIGHT, NAVY)
    for left, right in zip(host_boxes, host_boxes[1:]):
        draw_arrow(draw, (left[0][2], 920), (right[0][0], 920), NAVY, 6)

    # PIO slave and Linux SPI master are directly aligned; no line crosses another block.
    draw_polyline_arrow(draw, [(197, spi_slave[3]), (197, host_boxes[0][0][1])], color=CYAN, width=9, arrow_start=True)
    draw_label(draw, (520, 700), "SPI PIO-frame · MOSI / MISO / SCK / CSn", color=CYAN, font_size=20)

    return save_diagram(image, "05_flujo_interno_sniffer_host.png")


def create_network_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        980,
        "Red y servicios",
        "Separación entre el transporte local y la interfaz expuesta a la notebook",
    )
    pi_boundary = (60, 175, 1160, 880)
    draw.rounded_rectangle(pi_boundary, radius=24, fill="#FFFFFF", outline=f"#{NAVY}", width=4)
    draw.text((90, 195), "RASPBERRY PI 3B+", font=pil_font(27, True), fill=f"#{NAVY}")

    bridge = (120, 290, 460, 500)
    flask = (700, 290, 1040, 500)
    eth0 = (700, 640, 1040, 820)
    notebook = (1340, 580, 1710, 850)

    draw_box(draw, bridge, "Bridge SPI / HTTP", "127.0.0.1:5100\nSolo loopback\nNo expuesto", LIGHT, NAVY)
    draw_box(draw, flask, "Dashboard Flask", "0.0.0.0:5000\nConsume bridge local\nExpone /metrics", PALE_BLUE, BLUE)
    draw_box(draw, eth0, "eth0", "192.168.50.2/24\nEnlace directo", PALE_GOLD, GOLD)
    draw_box(draw, notebook, "Notebook", "192.168.50.1/24\nDashboard :5000\nPrometheus :9090\nGrafana :3000", PALE_GREEN, GREEN)

    draw_polyline_arrow(draw, [(bridge[2], 395), (flask[0], 395)], color=NAVY, width=8)
    draw_label(draw, ((bridge[2] + flask[0]) // 2, 350), "HTTP sobre 127.0.0.1", color=NAVY, font_size=20)
    draw_polyline_arrow(draw, [(870, flask[3]), (870, eth0[1])], color=GOLD, width=7, arrow_start=True)
    draw_label(draw, (1000, 570), "Servicio publicado en :5000", color=GOLD, font_size=19)
    draw_polyline_arrow(draw, [(eth0[2], 730), (notebook[0], 730)], color=CYAN, width=9, arrow_start=True)
    draw_label(draw, (1270, 680), "Ethernet directo", color=CYAN, font_size=19)

    draw.text((120, 560), "Criterio de exposición", font=pil_font(25, True), fill=f"#{NAVY}")
    note = (
        "La notebook no accede al bridge. Flask actúa como frontera de aplicación y Prometheus "
        "consulta las métricas publicadas por Flask mediante Ethernet."
    )
    draw.rounded_rectangle((120, 610, 540, 825), radius=16, fill=f"#{LIGHT}", outline=f"#{NAVY}", width=3)
    draw_centered_text(draw, (145, 625, 515, 810), note, pil_font(21), INK, 7)
    return save_diagram(image, "06_red_y_servicios.png")


def create_results_diagram() -> Path:
    image, draw = new_canvas(
        1800,
        880,
        "Validación final de ocho horas",
        "Sesión 20260606_022347_corrida_final_stream_8mhz_8h",
    )
    values = [
        ("Recibidas", 58_273_540, NAVY),
        ("Aceptadas", 17_482_062, GREEN),
        ("Filtradas", 40_791_478, BLUE),
    ]
    max_value = max(v for _, v, _ in values)
    y = 210
    for label, value, color in values:
        draw.text((100, y + 12), label, font=pil_font(27, True), fill=f"#{INK}")
        x1, x2 = 360, 1620
        draw.rounded_rectangle((x1, y, x2, y + 70), radius=14, fill=f"#{LIGHT}")
        bar_x = x1 + int((x2 - x1) * value / max_value)
        draw.rounded_rectangle((x1, y, bar_x, y + 70), radius=14, fill=f"#{color}")
        draw.text((bar_x + 18, y + 15), f"{value:,}".replace(",", "."), font=pil_font(25, True), fill=f"#{INK}")
        y += 125
    metrics = [
        ("SPI operativos", "0"),
        ("Paridad", "0"),
        ("Overflow / drops", "0 / 0"),
        ("Balance", "0"),
        ("Veredicto", "PASS"),
    ]
    x = 100
    for label, value in metrics:
        draw.rounded_rectangle((x, 625, x + 300, 790), radius=16, fill=f"#{PALE_GREEN}", outline=f"#{GREEN}", width=3)
        draw_centered_text(draw, (x + 10, 640, x + 290, 700), label, pil_font(20, True), GREEN)
        draw_centered_text(draw, (x + 10, 700, x + 290, 775), value, pil_font(34, True), NAVY)
        x += 325
    return save_diagram(image, "07_resultados_corrida_final.png")


def set_run_font(
    run,
    size: float = 11,
    color: str = INK,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_cell_text(cell, text: str, bold: bool = False, color: str = INK, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], LIGHT)
        add_cell_text(table.rows[0].cells[idx], header, bold=True, color=NAVY, size=9.5)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            add_cell_text(row.cells[idx], value, size=9.2)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(4)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def set_document_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def set_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("PAMPA / ARINC 429  |  Descripción del sistema")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "D7DEE6")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = ""
    footer_p.paragraph_format.space_before = Pt(0)
    add_page_number(footer_p)


def add_paragraph(
    doc: Document,
    text: str,
    *,
    bold_prefix: str | None = None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size: float = 11,
    color: str = INK,
    after: float = 6,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, size=size, color=color, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=size, color=color)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, color=color)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(3)
    r1 = p1.add_run(title)
    set_run_font(r1, size=10.5, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc: Document, path: Path, width: float, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_section_heading(doc: Document, number: str, title: str, intro: str | None = None) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run(f"{number}. {title}")
    set_run_font(run, size=16, color=BLUE, bold=True)
    if intro:
        add_paragraph(doc, intro, size=10.5, color=MUTED, after=8)


def build_document(diagrams: dict[str, Path]) -> None:
    doc = Document()
    set_document_styles(doc)
    for section in doc.sections:
        configure_section(section)
        set_header_footer(section)

    # Portada técnica y resumen.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("SISTEMA PAMPA / ARINC 429")
    set_run_font(r, size=12, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Descripción de arquitectura y validación de la fase lógica")
    set_run_font(r, size=27, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Transmisión, captura pasiva, transporte SPI y observabilidad")
    set_run_font(r, size=14, color=MUTED)

    add_table(
        doc,
        ["Versión", "Rama", "Estado", "Fecha"],
        [["arinc-logic-v1.0.0", "ARINC", "Fase lógica validada", "Junio de 2026"]],
        [2340, 1800, 2880, 2340],
    )
    add_picture(doc, diagrams["architecture"], 6.35, "Vista resumida del recorrido completo de los datos.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Resumen")
    set_run_font(r, size=13, color=BLUE, bold=True)
    add_paragraph(
        doc,
        "El sistema reproduce en laboratorio la estructura lógica y temporal de ARINC 429 mediante "
        "tres Raspberry Pi Pico. Un transmisor genera palabras sobre un enlace simplex FWD, un receptor "
        "las procesa y un tercer nodo actúa como observador pasivo. La captura se entrega por SPI a una "
        "Raspberry Pi 3B+, donde se publica mediante un bridge local, un dashboard Flask y métricas para "
        "Prometheus/Grafana.",
        size=10.2,
        after=5,
    )
    add_callout(
        doc,
        "Resultado principal",
        "La configuración final operó ocho horas a 8 MHz sobre SPI PIO-frame, procesó 58.273.540 "
        "palabras y mantuvo en cero los errores SPI operativos, errores de paridad, overflow, drops y "
        "resincronizaciones FWD operativas.",
        PALE_GREEN,
        GREEN,
    )

    page_break(doc)

    # Arquitectura general.
    add_section_heading(
        doc,
        "1",
        "Alcance y arquitectura general",
        "La implementación actual corresponde a una fase lógica y temporal. No representa todavía los niveles eléctricos de un bus ARINC 429 de campo.",
    )
    add_picture(doc, diagrams["architecture"], 6.45, "Bloques funcionales y enlaces principales.")
    add_table(
        doc,
        ["Componente", "Función principal", "Interfaz"],
        [
            ["ARINC-TX", "Genera palabras ARINC lógicas en flujo continuo y ráfagas variables.", "FWD GP2/GP3"],
            ["ARINC-RX", "Reconstruye y procesa palabras FWD. En stream no transmite ACK.", "FWD GP2/GP3"],
            ["ARINC-SNIFFER", "Toma pasiva sobre FWD/REV, valida paridad, filtra y mantiene snapshots.", "Derivación en cable + SPI"],
            ["Raspberry Pi 3B+", "Extrae datos por SPI y ejecuta bridge, Flask y registro.", "SPI + Ethernet"],
            ["Notebook", "Visualización, histórico y análisis.", "Ethernet directo"],
        ],
        [1800, 4860, 2700],
    )
    add_callout(
        doc,
        "Pasividad del SNIFFER",
        "El SNIFFER se conecta en paralelo sobre los conductores FWD/REV; no forma un enlace punto a punto "
        "con TX o RX y nunca transmite sobre esas líneas. Su única salida digital es la respuesta SPI hacia "
        "la Raspberry Pi 3B+.",
        PALE_GOLD,
        GOLD,
    )

    page_break(doc)

    # Conexionado.
    add_section_heading(
        doc,
        "2",
        "Conexionado físico",
        "TX y RX están unidos por el cable FWD. El SNIFFER no se conecta a una placa específica: sus entradas se derivan en paralelo desde ese mismo par de conductores.",
    )
    add_picture(doc, diagrams["wiring"], 6.45, "Esquema de conexión lógica, SPI y Ethernet.")
    add_table(
        doc,
        ["Enlace", "Origen", "Destino", "Señales"],
        [
            ["FWD", "ARINC-TX", "ARINC-RX", "FWD_A/FWD_B; SNIFFER derivado en paralelo"],
            ["REV opcional", "Otro canal simplex", "Receptor de ese canal", "REV_A/REV_B; SNIFFER derivado en paralelo"],
            ["SPI", "Raspberry Pi 3B+", "ARINC-SNIFFER", "MOSI GP16, CS GP17, SCK GP18, MISO GP19"],
            ["Ethernet", "Notebook", "Raspberry Pi 3B+", "192.168.50.1 ↔ 192.168.50.2"],
        ],
        [1260, 2160, 2520, 3420],
    )
    add_callout(
        doc,
        "Advertencia",
        "Los GPIO de 3,3 V no se conectan directamente a una línea ARINC real. La siguiente fase requiere "
        "protección, entrada de alta impedancia y un receptor ARINC 429 dedicado.",
        PALE_GOLD,
        RED,
    )

    page_break(doc)

    # Modelo lógico.
    add_section_heading(doc, "3", "Palabra y codificación ARINC lógica")
    add_picture(doc, diagrams["arinc"], 6.45, "Estructura de palabra y temporización bipolar RZ empleada.")
    add_table(
        doc,
        ["Propiedad", "Implementación"],
        [
            ["Velocidad", "100 kbps; cada bit ocupa 10 µs."],
            ["Retorno a cero", "5 µs en nivel activo y 5 µs en nivel NULL."],
            ["Representación diferencial", "Aproximadamente +3,3 V, 0 V y -3,3 V en CH1 - CH2."],
            ["Palabra", "32 bits: Label, SDI, datos, SSM y paridad impar."],
            ["Aceptación", "Paridad correcta, combinación Label/SDI habilitada y datos plausibles."],
        ],
        [2340, 7020],
    )
    add_paragraph(
        doc,
        "Las combinaciones operativas principales son 0xA5/0 para temperatura, 0xB1/1 para velocidad "
        "y 0xC2/2 para altitud. El modo histórico de laboratorio también utilizó 0xAC/3 para ACK_BATCH.",
        size=10.2,
    )

    page_break(doc)

    # TX/RX.
    add_section_heading(doc, "4", "Funcionamiento interno de TX y RX")
    add_picture(doc, diagrams["tx_rx"], 6.45, "Movimiento de la palabra desde la aplicación hasta los GPIO y regreso a la CPU receptora.")
    add_table(
        doc,
        ["Bloque", "Comportamiento relevante"],
        [
            ["Generador TX", "Construye palabras, alterna estados SSM y forma ráfagas de 1 a 2000 palabras."],
            ["PIO TX", "Serializa sin depender del tiempo de ejecución de la CPU."],
            ["PIO RX", "Muestrea el par lógico y reconstruye palabras de 32 bits."],
            ["FIFO RX", "Amortigua diferencias breves; conserva hasta cuatro palabras."],
            ["Lazo de captura", "Drena el FIFO de forma continua y evita trabajo prolongado en la ruta crítica."],
        ],
        [2340, 7020],
    )
    add_callout(
        doc,
        "Modo stream",
        "El flujo no depende de bloques fijos de 1000 palabras. La cantidad de tramas y los intervalos "
        "entre ráfagas varían; RX recibe en forma continua y REV permanece inactivo.",
        PALE_GREEN,
        GREEN,
    )

    page_break(doc)

    # Sniffer/host.
    add_section_heading(doc, "5", "Funcionamiento interno del SNIFFER y la 3B+")
    add_picture(doc, diagrams["sniffer_host"], 6.45, "Ruta de captura, clasificación y publicación.")
    add_table(
        doc,
        ["Etapa", "Resultado"],
        [
            ["Captura PIO", "Dos receptores independientes observan FWD y REV."],
            ["Paridad estricta", "Una palabra inválida incrementa el error, pero no actualiza accepted_words ni snapshot."],
            ["Filtro", "Whitelist de hasta 10 combinaciones Label/SDI."],
            ["Snapshot", "Hasta 16 variables recientes, además de contadores acumulados."],
            ["SPI PIO-frame", "Solicitud y respuesta en tramas completas de 32 bytes."],
            ["Bridge lock", "Serializa polling, filtros y reset para impedir transacciones SPI intercaladas."],
        ],
        [2340, 7020],
    )

    page_break(doc)

    # Network.
    add_section_heading(doc, "6", "Red, direcciones IP y servicios")
    add_picture(doc, diagrams["network"], 6.45, "Distribución de servicios entre la Raspberry Pi y la notebook.")
    add_table(
        doc,
        ["Elemento", "Dirección", "Uso"],
        [
            ["Notebook Ethernet", "192.168.50.1/24", "Acceso al dashboard y ejecución de observabilidad."],
            ["Raspberry Pi eth0", "192.168.50.2/24", "Servidor visible en la red directa."],
            ["Bridge", "127.0.0.1:5100", "API local para SPI; no se expone a la notebook."],
            ["Flask", "0.0.0.0:5000", "Dashboard en http://192.168.50.2:5000."],
            ["Prometheus", "Notebook :9090", "Consulta http://192.168.50.2:5000/metrics."],
            ["Grafana", "Notebook :3000", "Paneles e histórico sobre Prometheus."],
        ],
        [2160, 2700, 4500],
    )
    add_paragraph(
        doc,
        "Flask consulta al bridge mediante HTTP sobre la pila TCP/IP local de la 3B+ y la dirección "
        "127.0.0.1. Esta separación evita exponer el control SPI y deja una única interfaz de aplicación "
        "visible desde la notebook.",
        size=10.2,
    )

    page_break(doc)

    # Oscilloscope FWD.
    add_section_heading(doc, "7", "Validación por osciloscopio: canal FWD")
    add_picture(doc, FWD_SINGLE, 6.25, "GP2 respecto de masa: amplitud aproximada de 3,36 V y retorno a cero entre pulsos.")
    add_picture(doc, FWD_DIFF, 6.25, "MATH = CH1 - CH2: niveles aproximados de +3,36 V, 0 V y -3,44 V.")
    add_paragraph(
        doc,
        "La traza individual confirma el nivel GPIO y la media celda activa. La resta de ambos hilos "
        "muestra la bipolaridad lógica: los colores corresponden a las dos polaridades de la señal "
        "diferencial calculada, no a dos palabras superpuestas.",
        size=10.1,
    )

    page_break(doc)

    # Oscilloscope REV.
    add_section_heading(doc, "8", "Validación por osciloscopio: canal REV histórico")
    add_picture(doc, REV_SINGLE, 6.25, "GP4 en el modo batch/ACK: actividad agrupada y pulso activo cercano a 5 µs.")
    add_picture(doc, REV_DIFF, 6.25, "MATH del par REV: ráfaga con niveles positivos, nulos y negativos.")
    add_callout(
        doc,
        "Contexto de esta evidencia",
        "Estas capturas pertenecen al modo previo de laboratorio donde ARINC-RX emitía ACK. En el modo "
        "stream final REV no se utiliza para responder. El SNIFFER conserva sus entradas REV para observar "
        "un segundo canal simplex cuando exista.",
        PALE_GOLD,
        GOLD,
    )

    page_break(doc)

    # Results.
    add_section_heading(doc, "9", "Resultados de validación")
    add_picture(doc, diagrams["results"], 6.45, "Contadores relativos de la corrida final sin inyección de errores.")
    add_table(
        doc,
        ["Prueba", "Resultado", "Conclusión"],
        [
            ["Corrida final", "8,0004 h; 58.273.540 palabras; cero errores operativos.", "PASS"],
            ["Paridad estricta", "6.109 errores inyectados y descartados sobre 610.905 palabras.", "PASS"],
            ["Barrido SPI", "8 y 9 MHz PASS; 10 MHz parcial; 16 y 50 MHz FAIL.", "8 MHz seleccionado"],
            ["Stream aleatorio", "Ráfagas variables, sin lotes fijos y sin ACK.", "Validado"],
        ],
        [2160, 5400, 1800],
    )
    add_paragraph(
        doc,
        "En la corrida final, la clasificación conservó exactamente el total: 17.482.062 palabras "
        "aceptadas más 40.791.478 filtradas equivalen a las 58.273.540 recibidas. La distribución "
        "30 % / 70 % coincide con el perfil deliberado del transmisor.",
        size=10.2,
    )

    page_break(doc)

    # Operational profile and closing.
    add_section_heading(doc, "10", "Configuración recomendada y estado del desarrollo")
    add_table(
        doc,
        ["Parámetro", "Valor recomendado"],
        [
            ["Target TX", "arinc_tx_arinc429_logic_stream"],
            ["Target RX", "arinc_rx_arinc429_logic_stream"],
            ["Target SNIFFER", "sniffer_arinc429_logic_pio_frame"],
            ["SPI", "PIO-frame, 8 MHz, CS manual por GPIO8"],
            ["Temporización CS", "setup 100 µs, hold 100 µs, respuesta 2 ms"],
            ["Polling / stats", "6 ms / 60 ms"],
            ["Red", "Ethernet directo 192.168.50.1 ↔ 192.168.50.2"],
        ],
        [2700, 6660],
    )
    add_callout(
        doc,
        "Estado actual",
        "La fase lógica queda cerrada y reproducible: enlace FWD a 100 kbps, SNIFFER pasivo, paridad "
        "estricta, SPI por tramas completas a 8 MHz, exportación estadística y observabilidad.",
        PALE_GREEN,
        GREEN,
    )
    add_callout(
        doc,
        "Límite vigente",
        "La implementación usa niveles GPIO de 3,3 V y masa común. Todavía no incorpora aislamiento, "
        "protección de línea ni receptor diferencial compatible con los niveles de tensión y la impedancia "
        "de una instalación ARINC 429 real.",
        PALE_GOLD,
        RED,
    )
    add_paragraph(
        doc,
        "La siguiente fase debe comenzar por la especificación del frente eléctrico: conector, protección "
        "contra transitorios, entrada de alta impedancia, receptor ARINC dedicado, fuente de prueba y "
        "validación instrumental. Una vez establecida esa interfaz, la lógica de captura, filtrado, SPI y "
        "visualización desarrollada en esta fase puede reutilizarse.",
        size=10.5,
        after=10,
    )
    add_paragraph(
        doc,
        "Fuentes internas: PROJECT_CONTEXT.md, documentación de evidencias, diagramas de bloques y "
        "capturas almacenadas en imagenes/.",
        size=9,
        color=MUTED,
        after=0,
    )

    doc.save(DOCX_PATH)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    diagrams = {
        "architecture": create_architecture_diagram(),
        "wiring": create_wiring_diagram(),
        "arinc": create_arinc_model_diagram(),
        "tx_rx": create_internal_tx_rx_diagram(),
        "sniffer_host": create_internal_sniffer_host_diagram(),
        "network": create_network_diagram(),
        "results": create_results_diagram(),
    }
    build_document(diagrams)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
